"""DOCX-рендер финального отчёта (через python-docx, 6 секций)."""

from __future__ import annotations

from pathlib import Path
from typing import Any


__all__ = ["write_docx"]


def write_docx(data: dict[str, Any], path: str | Path) -> None:
    """Записать отчёт в .docx (через python-docx).

    Args:
        data: dict, возвращённый ``modes.synthesize.run()``.
        path: путь к выходному .docx-файлу.

    Raises:
        RuntimeError: если python-docx не установлен.
        OSError: если не удалось записать файл.
    """
    try:
        from docx import Document
    except ImportError as exc:
        raise RuntimeError(
            "python-docx не установлен; нельзя сохранить в .docx. "
            "Установите: pip install python-docx"
        ) from exc

    doc = Document()
    title = str(data.get("title") or "Анализ отклонения").strip() or "Анализ отклонения"
    doc.add_heading(title, level=1)

    date_iso = str(data.get("date_iso") or "")
    severity = str(data.get("severity") or "средняя")
    citations = data.get("vnd_citations") or []
    findings_count = int(data.get("source_findings_count") or len(citations))
    dropped = int(data.get("citations_dropped") or 0)
    normalized = str(data.get("normalized_violation") or "").strip()

    p = doc.add_paragraph()
    if date_iso:
        p.add_run(f"Дата: {date_iso}").bold = True
        p.add_run("  |  ")
    p.add_run(f"Категория тяжести: {severity}").bold = (date_iso == "")

    p2 = doc.add_paragraph()
    p2.add_run(
        f"Валидированных доказательств: {len(citations)} "
        f"(из {findings_count} кандидатов)"
    )

    if dropped:
        p3 = doc.add_paragraph()
        p3.add_run(
            f"Отклонено проверкой: {dropped} цитат не прошли сверку с источником"
        )

    if normalized:
        doc.add_paragraph(f"Нормализованная формулировка: {normalized}", style="Intense Quote")

    # 1. Краткое изложение отклонения
    doc.add_heading("1. Краткое изложение отклонения", level=2)
    summary = data.get("violation_summary") or []
    if summary:
        for line in summary:
            doc.add_paragraph(str(line))
    else:
        doc.add_paragraph("(секция не заполнена: LLM не вернул violation_summary)")

    # 2. Установленные факты (по ВНД)
    doc.add_heading("2. Установленные факты (по ВНД)", level=2)
    facts = data.get("established_facts") or []
    if facts:
        for line in facts:
            doc.add_paragraph(str(line))
    else:
        doc.add_paragraph(
            "(секция не заполнена: релевантных фрагментов ВНД не найдено)"
        )

    # 3. Анализ отклонения
    doc.add_heading("3. Анализ отклонения", level=2)
    analysis = data.get("deviation_analysis") or []
    if analysis:
        for line in analysis:
            doc.add_paragraph(str(line))
    else:
        doc.add_paragraph("(секция не заполнена: недостаточно данных для анализа)")

    # 4. Релевантные фрагменты ВНД (валидированные цитаты)
    doc.add_heading("4. Релевантные фрагменты ВНД (валидированные цитаты)", level=2)
    if citations:
        for i, c in enumerate(citations, 1):
            evidence_id = str(c.get("evidence_id") or f"F{i}")
            source = str(c.get("source_file") or "(неизвестный источник)")
            chunk_index = c.get("chunk_index")
            relation = str(c.get("relation_type") or "контекст")
            excerpt = str(c.get("excerpt") or "").strip()
            explanation = str(c.get("relation_explanation") or "").strip()

            heading = f"{i}. [{evidence_id}] {source}"
            if chunk_index is not None:
                heading += f" — чанк {chunk_index}"
            doc.add_heading(heading, level=3)

            meta = doc.add_paragraph()
            meta.add_run("Тип соотнесения: ").bold = True
            meta.add_run(relation)

            if excerpt:
                doc.add_paragraph(f"«{excerpt}»", style="Intense Quote")
            if explanation:
                doc.add_paragraph(explanation)
    else:
        if findings_count == 0:
            doc.add_paragraph(
                "(секция не заполнена: релевантных фрагментов ВНД не найдено)"
            )
        else:
            doc.add_paragraph(
                f"(секция не заполнена: все {findings_count} цитат отклонены проверкой)"
            )

    # 5. Итоговая классификация
    doc.add_heading("5. Итоговая классификация", level=2)
    verdict = data.get("verdict") or {}
    category = str(verdict.get("category") or severity or "средняя")
    p = doc.add_paragraph()
    p.add_run("Категория: ").bold = True
    p.add_run(category)
    verdict_text = verdict.get("verdict_text") or []
    if verdict_text:
        for line in verdict_text:
            doc.add_paragraph(str(line))
    else:
        doc.add_paragraph("(секция не заполнена: LLM не вернул verdict_text)")

    # 6. Рекомендуемая усиленная формулировка
    doc.add_heading("6. Рекомендуемая усиленная формулировка", level=2)
    formulation = data.get("recommended_formulation") or []
    if formulation:
        for line in formulation:
            doc.add_paragraph(str(line))
    else:
        doc.add_paragraph(
            "(секция не заполнена: LLM не вернул recommended_formulation)"
        )

    doc.save(str(path))
